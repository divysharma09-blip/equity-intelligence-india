"""Excel research workbook (data/Equity_Intelligence_India.xlsx).

Update rules — previous research is never deleted:
* Master_Companies: one row per company, updated in place (matched on Company Key).
* Financial_History, Forecasts, Valuation, Factor_Scores, Risks_Governance,
  Sources, Research_History: rows are APPENDED per saved report (Report ID column).
* Dashboard and Settings_Methodology are derived summaries and are regenerated.

Each write backs the existing workbook up to data/backups/ first, then saves
atomically (temp file + rename) so a crash cannot leave a half-written file.
Columns are located by header text, so user-reordered columns stay intact.
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from eii.config import APP_NAME, FRAMEWORK_DISCLAIMER
from eii.db.database import copy_with_backup
from eii.models import DATA_STATUS_LABELS, DataStatus, ResearchReport
from eii.scoring import DEFAULT_SCORECARD, ScorecardConfig, score_band

log = logging.getLogger("eii.export")

PCT, NUM, INT, X = "0.0%", "#,##0.00", "#,##0", '0.0"x"'


@dataclass(frozen=True)
class Col:
    header: str
    width: int = 16
    fmt: Optional[str] = None


SHEETS: dict[str, list[Col]] = {
    "Dashboard": [],  # free-form, regenerated
    "Master_Companies": [
        Col("Company Key", 16), Col("Company Name", 36), Col("NSE Symbol", 14), Col("BSE Code", 11),
        Col("ISIN", 15), Col("Sector", 20), Col("Industry", 26), Col("Verification Status", 18),
        Col("First Saved (UTC)", 20), Col("Last Updated (UTC)", 20), Col("Latest Report ID", 18),
        Col("Latest Version", 9, INT), Col("Framework Score", 10, "0.0"), Col("Score Coverage", 10, PCT),
        Col("Conclusion Allowed", 12), Col("Company Type", 12), Col("Statement Basis", 13)],
    "Financial_History": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Metric", 24), Col("Fiscal Period End", 14),
        Col("Value", 20, NUM), Col("Unit", 10), Col("Data Status", 16), Col("Source ID", 16),
        Col("Saved At (UTC)", 20)],
    "Forecasts": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Scenario", 9), Col("Horizon (yrs)", 9, INT),
        Col("Year", 6, INT), Col("Projected EPS", 14, NUM), Col("Data Status", 26), Col("Basis", 60),
        Col("Saved At (UTC)", 20)],
    "Valuation": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Scenario", 9), Col("Horizon (yrs)", 9, INT),
        Col("Start Price", 12, NUM), Col("Start EPS", 10, NUM), Col("EPS Growth (assumed)", 12, PCT),
        Col("Exit P/E (assumed)", 11, "0.0"), Col("Dividend Yield (assumed)", 12, PCT),
        Col("Target Price", 12, NUM), Col("Total Return", 11, PCT), Col("Expected CAGR", 11, PCT),
        Col("Rationale", 60), Col("Data Status", 12), Col("Saved At (UTC)", 20)],
    "Factor_Scores": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Factor", 24), Col("Score (0-100)", 10, "0.0"),
        Col("Weight", 8, "0.0"), Col("Coverage", 9, PCT), Col("Metric Details", 80),
        Col("Saved At (UTC)", 20)],
    "Risks_Governance": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Type", 11), Col("Category / Item", 24),
        Col("Description / Value", 60), Col("Severity", 9), Col("Mitigation / Notes", 40),
        Col("Data Status", 14), Col("Saved At (UTC)", 20)],
    "Sources": [
        Col("Company Key", 16), Col("Report ID", 18), Col("Source ID", 16), Col("Title", 50),
        Col("Provider", 30), Col("Kind", 20), Col("URL", 50), Col("Retrieved At (UTC)", 20),
        Col("Notes", 60)],
    "Research_History": [
        Col("Saved At (UTC)", 20), Col("Company Key", 16), Col("Company Name", 36), Col("Report ID", 18),
        Col("Version", 8, INT), Col("Status", 10), Col("Verification", 18), Col("Framework Score", 10, "0.0"),
        Col("Score Coverage", 10, PCT), Col("Conclusion Allowed", 12), Col("Analyst", 14), Col("Note", 60)],
    "Settings_Methodology": [],  # free-form, regenerated
}
SHEET_ORDER = list(SHEETS)

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


class ExcelExportError(RuntimeError):
    """User-readable Excel failure."""


def _safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
    return value


def _set(ws: Worksheet, row: int, col: int, value: Any, fmt: Optional[str] = None) -> None:
    cell = ws.cell(row=row, column=col, value=_safe(value))
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # never let external text become a formula
    if fmt and value is not None:
        cell.number_format = fmt


def _write_header(ws: Worksheet, cols: list[Col]) -> None:
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c.header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = c.width
    ws.freeze_panes = "A2"


def _header_map(ws: Worksheet, cols: list[Col]) -> dict[str, int]:
    """Header -> column index; appends any expected header missing from the sheet."""
    existing = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    next_col = max(existing.values(), default=0) + 1
    for c in cols:
        if c.header not in existing:
            cell = ws.cell(row=1, column=next_col, value=c.header)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            ws.column_dimensions[get_column_letter(next_col)].width = c.width
            existing[c.header] = next_col
            next_col += 1
    return existing


def _append(ws: Worksheet, cols: list[Col], record: dict[str, Any]) -> None:
    hmap = _header_map(ws, cols)
    fmts = {c.header: c.fmt for c in cols}
    row = ws.max_row + 1
    for header, value in record.items():
        _set(ws, row, hmap[header], value, fmts.get(header))


def _iso(dt: Optional[datetime]) -> Optional[str]:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S") if dt else None


class WorkbookManager:
    def __init__(self, path: Path, backup_dir: Path, max_backups: int = 20,
                 config: ScorecardConfig = DEFAULT_SCORECARD):
        self.path = path
        self.backup_dir = backup_dir
        self.max_backups = max_backups
        self.config = config

    # ---------------------------------------------------------------- public

    def ensure_workbook(self) -> None:
        """Create the workbook with all sheets if it does not exist."""
        if self.path.exists():
            return
        wb = Workbook()
        wb.remove(wb.active)
        self._ensure_sheets(wb)
        self._render_dashboard(wb)
        self._render_methodology(wb)
        self._atomic_save(wb)
        log.info("Created workbook %s", self.path)

    def append_report(self, report: ResearchReport, version: int) -> None:
        """Add a saved report. Existing companies and history rows are preserved."""
        self._check_not_open()
        wb = self._load()
        copy_with_backup(self.path, self.backup_dir, self.max_backups)
        self._ensure_sheets(wb)
        saved_at = _iso(report.updated_at)
        key = report.identity.key
        rid = report.report_id
        self._upsert_master(wb["Master_Companies"], report, version, saved_at)

        fh = wb["Financial_History"]
        for f in sorted(report.financials, key=lambda f: (f.metric, f.period_end)):
            _append(fh, SHEETS["Financial_History"], {
                "Company Key": key, "Report ID": rid, "Metric": f.metric,
                "Fiscal Period End": f.period_end.isoformat(), "Value": f.value, "Unit": f.unit,
                "Data Status": DATA_STATUS_LABELS[f.status], "Source ID": f.source_id,
                "Saved At (UTC)": saved_at})

        assumptions = {(a.scenario, a.horizon_years): a for a in report.scenarios}
        for r in report.scenario_results:
            a = assumptions.get((r.scenario, r.horizon_years))
            basis = (f"EPS {r.start_eps:,.2f} grown at assumed {a.eps_growth:.1%}/yr" if a else "")
            for year, eps in enumerate(r.eps_path, start=1):
                _append(wb["Forecasts"], SHEETS["Forecasts"], {
                    "Company Key": key, "Report ID": rid, "Scenario": r.scenario,
                    "Horizon (yrs)": r.horizon_years, "Year": year, "Projected EPS": eps,
                    "Data Status": "Assumption-driven projection (not a forecast)", "Basis": basis,
                    "Saved At (UTC)": saved_at})
            _append(wb["Valuation"], SHEETS["Valuation"], {
                "Company Key": key, "Report ID": rid, "Scenario": r.scenario,
                "Horizon (yrs)": r.horizon_years, "Start Price": r.start_price, "Start EPS": r.start_eps,
                "EPS Growth (assumed)": a.eps_growth if a else None,
                "Exit P/E (assumed)": a.exit_pe if a else None,
                "Dividend Yield (assumed)": a.dividend_yield if a else None,
                "Target Price": r.target_price, "Total Return": r.total_return, "Expected CAGR": r.cagr,
                "Rationale": a.rationale if a else "", "Data Status": DATA_STATUS_LABELS[DataStatus.ASSUMPTION],
                "Saved At (UTC)": saved_at})

        if report.scorecard:
            for fs in report.scorecard.factors:
                details = "; ".join(
                    f"{m.label}: " + ("n/a" if m.value is None else f"{m.value:.4g}")
                    + ("" if m.score is None else f" → {m.score:.0f}")
                    + (f" [{m.excluded_reason}]" if m.excluded_reason else "") for m in fs.metrics)
                _append(wb["Factor_Scores"], SHEETS["Factor_Scores"], {
                    "Company Key": key, "Report ID": rid, "Factor": fs.label, "Score (0-100)": fs.score,
                    "Weight": fs.weight, "Coverage": fs.coverage, "Metric Details": details,
                    "Saved At (UTC)": saved_at})

        rg = wb["Risks_Governance"]
        for risk in report.risks:
            _append(rg, SHEETS["Risks_Governance"], {
                "Company Key": key, "Report ID": rid, "Type": "Risk", "Category / Item": risk.category,
                "Description / Value": risk.description, "Severity": risk.severity,
                "Mitigation / Notes": risk.mitigation, "Data Status": DATA_STATUS_LABELS[risk.status],
                "Saved At (UTC)": saved_at})
        g = report.governance
        for item, value in (("Promoter holding %", g.promoter_holding_pct),
                            ("Promoter pledge %", g.promoter_pledge_pct), ("Auditor", g.auditor or None),
                            ("Governance rating (1-5)", g.governance_rating),
                            ("Macro resilience rating (1-5)", g.macro_resilience_rating),
                            ("Governance notes", g.notes or None)):
            if value is not None:
                _append(rg, SHEETS["Risks_Governance"], {
                    "Company Key": key, "Report ID": rid, "Type": "Governance", "Category / Item": item,
                    "Description / Value": str(value), "Data Status": DATA_STATUS_LABELS[DataStatus.USER_ENTERED],
                    "Saved At (UTC)": saved_at})

        own, aud = report.ownership, report.audit
        filed: list[tuple[str, object, str]] = []
        if own:
            filed += [("Promoter holding % (filed)", own.promoter_holding_pct, f"as of {own.as_of}"),
                      ("Promoter shares pledged % of promoter holding (filed)", own.promoter_pledged_pct_of_promoter,
                       f"as of {own.as_of}"),
                      ("Pledged promoter shares % of total (filed)", own.promoter_pledged_pct_of_total, f"as of {own.as_of}")]
        if aud:
            filed += [("Statutory auditor(s) (filed)", ", ".join(aud.auditors) or None, f"FY{aud.period_end.year}"),
                      ("Audit opinion (filed)", aud.opinion, f"FY{aud.period_end.year}")]
        for item, value, note in filed:
            if value is not None:
                _append(rg, SHEETS["Risks_Governance"], {
                    "Company Key": key, "Report ID": rid, "Type": "Governance", "Category / Item": item,
                    "Description / Value": str(value), "Mitigation / Notes": note,
                    "Data Status": DATA_STATUS_LABELS[DataStatus.REPORTED], "Saved At (UTC)": saved_at})

        for s in report.sources:
            _append(wb["Sources"], SHEETS["Sources"], {
                "Company Key": key, "Report ID": rid, "Source ID": s.source_id, "Title": s.title,
                "Provider": s.provider, "Kind": s.kind.value, "URL": s.url,
                "Retrieved At (UTC)": _iso(s.retrieved_at), "Notes": s.notes})

        sc = report.scorecard
        _append(wb["Research_History"], SHEETS["Research_History"], {
            "Saved At (UTC)": saved_at, "Company Key": key, "Company Name": report.identity.company_name,
            "Report ID": rid, "Version": version, "Status": report.status,
            "Verification": report.verification.status.value if report.verification else "unverified",
            "Framework Score": sc.overall_score if sc else None, "Score Coverage": sc.coverage if sc else None,
            "Conclusion Allowed": "Yes" if report.readiness and report.readiness.can_conclude else "No",
            "Analyst": report.analyst,
            "Note": f"{len(report.financials)} facts, {len(report.sources)} sources, "
                    f"{len(report.risks)} risks, {len(report.warnings)} warnings"})

        self._render_dashboard(wb)
        self._render_methodology(wb)
        self._atomic_save(wb)
        log.info("Workbook updated with report %s (%s v%s)", rid, key, version)

    # ---------------------------------------------------------------- internals

    def _check_not_open(self) -> None:
        lock = self.path.with_name(f"~${self.path.name}")
        if lock.exists():
            raise ExcelExportError(f"The workbook appears to be open in Excel ({lock.name} exists). "
                                   "Close it and retry the Excel sync.")

    def _load(self):
        if not self.path.exists():
            self.ensure_workbook()
        try:
            return load_workbook(self.path)
        except Exception as exc:
            raise ExcelExportError(f"Could not open {self.path.name}: {exc}. A backup of each previous "
                                   f"version is in {self.backup_dir}.") from exc

    def _ensure_sheets(self, wb: Workbook) -> None:
        for name in SHEET_ORDER:
            if name not in wb.sheetnames:
                ws = wb.create_sheet(name)
                if SHEETS[name]:
                    _write_header(ws, SHEETS[name])
            elif SHEETS[name]:
                _header_map(wb[name], SHEETS[name])
        # keep the canonical sheets first and in order; user-added sheets follow
        wb._sheets.sort(key=lambda s: SHEET_ORDER.index(s.title) if s.title in SHEET_ORDER else 99)

    def _upsert_master(self, ws: Worksheet, report: ResearchReport, version: int, saved_at: str) -> None:
        cols = SHEETS["Master_Companies"]
        hmap = _header_map(ws, cols)
        key = report.identity.key
        key_col = hmap["Company Key"]
        row = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, key_col).value == key), None)
        first_saved = ws.cell(row, hmap["First Saved (UTC)"]).value if row else saved_at
        row = row or ws.max_row + 1
        ident, prof, sc = report.identity, report.profile, report.scorecard
        record = {
            "Company Key": key, "Company Name": ident.company_name, "NSE Symbol": ident.nse_symbol,
            "BSE Code": ident.bse_code, "ISIN": ident.isin,
            "Sector": prof.sector if prof else None, "Industry": prof.industry if prof else None,
            "Verification Status": report.verification.status.value if report.verification else "unverified",
            "First Saved (UTC)": first_saved, "Last Updated (UTC)": saved_at,
            "Latest Report ID": report.report_id, "Latest Version": version,
            "Framework Score": sc.overall_score if sc else None, "Score Coverage": sc.coverage if sc else None,
            "Conclusion Allowed": "Yes" if report.readiness and report.readiness.can_conclude else "No",
            "Company Type": report.company_type, "Statement Basis": report.statement_basis,
        }
        fmts = {c.header: c.fmt for c in cols}
        for header, value in record.items():
            if value is None and header in ("NSE Symbol", "BSE Code", "ISIN", "Sector", "Industry"):
                continue  # keep previously known values
            _set(ws, row, hmap[header], value, fmts.get(header))

    def _render_dashboard(self, wb: Workbook) -> None:
        ws = wb["Dashboard"]
        ws.delete_rows(1, ws.max_row)
        ws["A1"], ws["A1"].font = APP_NAME + " — Research Dashboard", TITLE_FONT
        ws["A2"] = f"Regenerated {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} from Master_Companies."
        ws["A3"] = FRAMEWORK_DISCLAIMER
        ws["A4"] = ("Scores are only meaningful where 'Conclusion Allowed' is Yes (verified identity, "
                    "≥3 years of data, ≥60% coverage). This workbook contains no buy/sell recommendations.")
        master = wb["Master_Companies"]
        hmap = {str(c.value): c.column for c in master[1] if c.value}
        companies = [{h: master.cell(r, c).value for h, c in hmap.items()} for r in range(2, master.max_row + 1)
                     if master.cell(r, hmap.get("Company Key", 1)).value]
        hist = wb["Research_History"]
        ws["A6"], ws["A6"].font = "Companies researched", BOLD
        ws["B6"] = len(companies)
        ws["A7"], ws["A7"].font = "Saved report versions", BOLD
        ws["B7"] = max(hist.max_row - 1, 0)
        headers = ["Company Name", "NSE Symbol", "ISIN", "Sector", "Verification Status", "Framework Score",
                   "Score Band", "Score Coverage", "Conclusion Allowed", "Last Updated (UTC)"]
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=9, column=i, value=h)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
        for r, comp in enumerate(sorted(companies, key=lambda c: str(c.get("Company Name") or "")), start=10):
            score = comp.get("Framework Score")
            allowed = comp.get("Conclusion Allowed") == "Yes"
            values = [comp.get("Company Name"), comp.get("NSE Symbol"), comp.get("ISIN"), comp.get("Sector"),
                      comp.get("Verification Status"), score if allowed else "Withheld",
                      score_band(score) if allowed else "Insufficient verified data — no conclusion",
                      comp.get("Score Coverage"), comp.get("Conclusion Allowed"), comp.get("Last Updated (UTC)")]
            for i, v in enumerate(values, start=1):
                _set(ws, r, i, v, {6: "0.0", 8: PCT}.get(i))
        for i, w in enumerate([36, 12, 15, 20, 18, 10, 40, 10, 12, 20], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _render_methodology(self, wb: Workbook) -> None:
        ws = wb["Settings_Methodology"]
        ws.delete_rows(1, ws.max_row)
        rows: list[tuple] = [
            (APP_NAME + " — Settings & Methodology",), (f"Scorecard version {self.config.version}",),
            (FRAMEWORK_DISCLAIMER,), (),
            ("DATA STATUS LABELS",), *[(s.value, label) for s, label in DATA_STATUS_LABELS.items()], (),
            ("SCORECARD", "Weight", "Metric", "Rule"),
        ]
        for f in self.config.factors:
            rows.append((f.label, f.weight, "", f.description))
            rows += [("", "", m.label, m.describe()) for m in f.metrics]
        rows += [
            (),
            ("SCORING RULES",),
            ("Metric score", "Linear between 'poor' (0) and 'excellent' (100), clipped."),
            ("Factor score", "Mean of available metric scores. Missing metrics are excluded, not zero-filled."),
            ("Overall score", "Weight-normalised mean of factors with data."),
            ("Coverage", "Weighted share of metrics that had data."),
            ("Conclusion gate", "Verified identity, ≥3 years statements, ≥60% coverage, current price, "
                                "no critical warnings. Otherwise no conclusion is shown."),
            ("Company types", "Each metric applies to industrial, bank and/or NBFC companies. Non-applicable "
                              "metrics are excluded from score and coverage. Insurers are not supported: no conclusion."),
            ("Filed vs analyst", "Filed promoter pledge and audit opinion take precedence over analyst entries."),
            (),
            ("DATA SOURCES",),
            ("Statements", "NSE XBRL annual results filings (legacy 2019–2024 and SEBI Integrated Filing 2025+). "
                           "One basis for all years: consolidated, or standalone for banks. Yahoo Finance is a "
                           "fallback and a cross-check (3% tolerance)."),
            ("Per-share data", "EPS and share counts restated for splits/bonuses from NSE corporate actions; "
                               "as-filed EPS kept as diluted_eps_reported."),
            ("Inconsistent filings", "If EPS × shares and net profit disagree by >20%, both are dropped for that year."),
            ("Ownership", "Promoter/public holding from the NSE shareholding pattern; pledge from NSE SAST "
                          "disclosures, verified against share counts."),
            ("Prices", "Yahoo Finance (unofficial); NSE's quote API blocks scripts."),
            (),
            ("RATIO DEFINITIONS (end-of-period balances)",),
            ("ROE", "Net profit / shareholders' equity (equity must be positive)"),
            ("ROCE", "EBIT / (total assets − current liabilities)"),
            ("Debt / equity", "Total debt / shareholders' equity"),
            ("Interest coverage", "EBIT / interest expense"),
            ("FCF", "Reported free cash flow; else operating cash flow − |capex| (derived)"),
            ("Cash conversion", "Σ operating cash flow / Σ net profit, last 3 years"),
            ("Accruals ratio", "(Net profit − operating cash flow) / total assets"),
            ("CAGR", "Longest span up to 3 fiscal years; not computed if either end ≤ 0"),
            ("EV", "Market cap + total debt − cash & short-term investments"),
            ("EBIT (filings)", "Profit before exceptional items and tax + finance costs (derived)"),
            ("Bank: NII", "Interest earned − interest expended (derived)"),
            ("Bank: cost-to-income", "Operating expenses / (NII + other income)"),
            ("Credit cost", "Provisions (or impairment) / year-end advances"),
            ("NII / assets", "Proxy for net interest margin (year-end assets, not average earning assets)"),
            (),
            ("SCENARIO MODEL",),
            ("Projection", "EPS_t = EPS_0 × (1+g)^t; target = EPS_n × exit P/E"),
            ("Expected CAGR", "(target/price)^(1/n) − 1 + assumed dividend yield"),
            ("Status", "All growth, P/E and yield inputs are ASSUMPTIONS, not forecasts."),
            (),
            ("IDENTITY VERIFICATION",),
            ("Rule", "Name match alone is never sufficient. Requires valid ISIN, presence in an official "
                     "exchange list with matching symbol↔ISIN, and matching data-provider name/ticker."),
        ]
        for r, values in enumerate(rows, start=1):
            for c, v in enumerate(values, start=1):
                _set(ws, r, c, v)
            if len(values) == 1 or (values and values[0] in ("SCORECARD",)):
                ws.cell(r, 1).font = BOLD
        ws["A1"].font = TITLE_FONT
        for col, w in zip("ABCD", (34, 60, 30, 60)):
            ws.column_dimensions[col].width = w

    def _atomic_save(self, wb: Workbook) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_name(f".{self.path.stem}.tmp.xlsx")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except OSError as exc:
            tmp.unlink(missing_ok=True)
            raise ExcelExportError(f"Could not write {self.path}: {exc}") from exc

