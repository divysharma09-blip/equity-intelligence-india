import os

import pytest
from openpyxl import load_workbook

from conftest import BETA_ISIN
from eii.export import SHEET_ORDER, ExcelExportError
from eii.models import (GovernanceInputs, ResearchReport, Risk, ScenarioAssumption, SecurityIdentity,
                        Source, SourceKind)
from eii.services.persistence import retry_excel_sync, save_completed_report
from eii.services.research import refresh_analysis

EXPECTED_SHEETS = ["Dashboard", "Master_Companies", "Financial_History", "Forecasts", "Valuation",
                   "Factor_Scores", "Risks_Governance", "Sources", "Research_History", "Settings_Methodology"]


@pytest.fixture
def alpha_report(alpha_identity, facts, market, profile, fin_source):
    return refresh_analysis(ResearchReport(
        identity=alpha_identity, financials=facts, market=market, profile=profile, sources=[fin_source],
        governance=GovernanceInputs(governance_rating=4, promoter_pledge_pct=0, auditor="Fictional & Co"),
        risks=[Risk(category="Regulatory", description="=HYPERLINK(\"http://evil\")", severity="medium")],
        scenarios=[ScenarioAssumption(scenario=s, horizon_years=3, eps_growth=g, exit_pe=pe)
                   for s, g, pe in (("bear", 0.0, 15), ("base", 0.1, 20), ("bull", 0.15, 25))]))


@pytest.fixture
def beta_report():
    ident = SecurityIdentity(company_name="Beta Finance Limited", nse_symbol="BETAFIN", isin=BETA_ISIN)
    return refresh_analysis(ResearchReport(identity=ident, sources=[Source(
        title="Analyst notes", provider="analyst", kind=SourceKind.USER_SUPPLIED)]))


def rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def col(ws, header):
    idx = [c.value for c in ws[1]].index(header)
    return [r[idx] for r in rows(ws)]


def test_requirement_sheet_names_match():
    assert SHEET_ORDER == EXPECTED_SHEETS


def test_ensure_workbook_creates_all_sheets(workbook, settings):
    workbook.ensure_workbook()
    wb = load_workbook(settings.excel_path)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb["Master_Companies"]["A1"].value == "Company Key"
    assert "not a scientifically validated" in str(wb["Settings_Methodology"]["A3"].value)


def test_save_appends_without_deleting_previous(repo, workbook, settings, alpha_report, beta_report):
    workbook.ensure_workbook()
    o1 = save_completed_report(alpha_report, repo, workbook)
    o2 = save_completed_report(beta_report, repo, workbook)
    o3 = save_completed_report(alpha_report.model_copy(update={"business_model": "v2"}), repo, workbook)
    assert o1.excel_ok and o2.excel_ok and o3.excel_ok, (o1.excel_error, o2.excel_error, o3.excel_error)
    assert o3.saved.version == 2

    wb = load_workbook(settings.excel_path)
    master = wb["Master_Companies"]
    assert sorted(col(master, "Company Key")) == sorted([alpha_report.identity.isin, BETA_ISIN])
    alpha_row = col(master, "Company Key").index(alpha_report.identity.isin)
    assert col(master, "Latest Version")[alpha_row] == 2
    assert col(master, "Latest Report ID")[alpha_row] == o3.saved.report_id

    hist = wb["Research_History"]
    assert col(hist, "Report ID") == [o1.saved.report_id, o2.saved.report_id, o3.saved.report_id]

    fh_ids = col(wb["Financial_History"], "Report ID")
    assert fh_ids.count(o1.saved.report_id) == len(alpha_report.financials)
    assert fh_ids.count(o3.saved.report_id) == len(alpha_report.financials)

    assert len(rows(wb["Valuation"])) == 6            # 3 scenarios × 2 alpha saves
    assert len(rows(wb["Forecasts"])) == 6 * 3         # 3 projected years each
    assert set(col(wb["Forecasts"], "Data Status")) == {"Assumption-driven projection (not a forecast)"}
    assert len(rows(wb["Factor_Scores"])) == 8 * 3
    assert col(wb["Sources"], "Report ID").count(o2.saved.report_id) == 1
    assert wb["Dashboard"]["B6"].value == 2 and wb["Dashboard"]["B7"].value == 3


def test_beta_report_without_data_is_marked_no_conclusion(repo, workbook, settings, beta_report):
    save_completed_report(beta_report, repo, workbook)
    wb = load_workbook(settings.excel_path)
    assert col(wb["Master_Companies"], "Conclusion Allowed") == ["No"]
    assert col(wb["Master_Companies"], "Framework Score") == [None]  # not fabricated
    dash = [r for r in wb["Dashboard"].iter_rows(min_row=10, values_only=True) if r[0]]
    assert dash[0][5] == "Withheld"
    assert dash[0][6] == "Insufficient verified data — no conclusion"


def test_user_edits_in_workbook_survive(repo, workbook, settings, alpha_report, beta_report):
    save_completed_report(alpha_report, repo, workbook)
    wb = load_workbook(settings.excel_path)
    wb.create_sheet("My_Notes")["A1"] = "keep me"
    hist = wb["Research_History"]
    hist.cell(hist.max_row + 1, 1, "manual row")
    wb.save(settings.excel_path)

    save_completed_report(beta_report, repo, workbook)
    wb = load_workbook(settings.excel_path)
    assert wb["My_Notes"]["A1"].value == "keep me"
    assert "manual row" in col(wb["Research_History"], "Saved At (UTC)")
    assert len(rows(wb["Research_History"])) == 3


def test_formula_injection_is_neutralised(repo, workbook, settings, alpha_report):
    save_completed_report(alpha_report, repo, workbook)
    ws = load_workbook(settings.excel_path)["Risks_Governance"]
    cell = next(c for row in ws.iter_rows(min_row=2) for c in row
                if isinstance(c.value, str) and c.value.startswith("=HYPERLINK"))
    assert cell.data_type == "s"


def test_backup_created_before_each_update(repo, workbook, settings, alpha_report):
    workbook.ensure_workbook()
    save_completed_report(alpha_report, repo, workbook)
    save_completed_report(alpha_report, repo, workbook)
    backups = list(settings.backup_dir.glob("Equity_Intelligence_India_*.xlsx"))
    assert len(backups) == 2


def test_open_workbook_keeps_db_save_and_allows_retry(repo, workbook, settings, alpha_report):
    workbook.ensure_workbook()
    lock = settings.excel_path.with_name(f"~${settings.excel_path.name}")
    lock.write_text("")
    outcome = save_completed_report(alpha_report, repo, workbook)
    assert not outcome.excel_ok and "open in Excel" in outcome.excel_error
    assert repo.get_report(outcome.saved.report_id) is not None   # SQLite save kept
    assert repo.unsynced_reports() == [outcome.saved.report_id]
    assert rows(load_workbook(settings.excel_path)["Research_History"]) == []

    os.remove(lock)
    synced, errors = retry_excel_sync(repo, workbook)
    assert synced == 1 and errors == []
    assert len(rows(load_workbook(settings.excel_path)["Research_History"])) == 1


def test_declining_save_modifies_nothing(repo, workbook, settings, alpha_report):
    """The 'No' path: drafting and recomputing a report must not touch either store."""
    workbook.ensure_workbook()
    before_xlsx = settings.excel_path.stat().st_mtime_ns
    before_db = settings.db_path.read_bytes()
    refresh_analysis(alpha_report.model_copy(update={"business_model": "draft only"}))
    assert settings.excel_path.stat().st_mtime_ns == before_xlsx
    assert settings.db_path.read_bytes() == before_db
    assert repo.list_companies().empty


def test_corrupt_workbook_gives_helpful_error(workbook, settings, alpha_report):
    settings.excel_path.write_text("not an xlsx")
    with pytest.raises(ExcelExportError, match="backup"):
        workbook.append_report(alpha_report, 1)
